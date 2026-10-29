from openpyxl import load_workbook

book = load_workbook("weekly_report.xlsx")
sheet = book["Courses"]
for row in sheet.iter_rows(values_only=True):
    print(row)
