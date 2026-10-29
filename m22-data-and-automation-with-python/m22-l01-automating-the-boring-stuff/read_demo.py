from pathlib import Path

from openpyxl import load_workbook

from sheets import summarise

path = Path("data/enrollments_week.xlsx")
ws = load_workbook(path).active
print(ws.title, ws.max_row, "rows,", ws.max_column, "columns")
for course, n, rupees in summarise(path):
    print(f"{course:24} {n}  {rupees}")
