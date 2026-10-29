from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


def summarise(path: Path) -> list[tuple[str, int, int]]:
    ws = load_workbook(path).active
    count, revenue = Counter(), Counter()
    rows = ws.iter_rows(min_row=2, values_only=True)
    for day, student, course, price in rows:
        count[course] += 1
        revenue[course] += price
    return [(c, count[c], rupees)
            for c, rupees in revenue.most_common()]


def write_report(summary, out: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Revenue"
    ws.append(["Course", "Enrollments", "Revenue", "Share"])
    for course, n, rupees in summary:
        ws.append([course, n, rupees])
    last = len(summary) + 1
    total = ["Total", f"=SUM(B2:B{last})", f"=SUM(C2:C{last})"]
    ws.append(total)
    for row in range(2, last + 1):
        ws[f"D{row}"] = f"=C{row}/C{last + 1}"
    for row in range(2, last + 2):
        ws[f"C{row}"].number_format = '"Rs "#,##0'
    for row in range(2, last + 1):
        ws[f"D{row}"].number_format = "0%"
    for cell in ws[1]:
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1A7CFF")
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 14
    ws.freeze_panes = "A2"
    wb.save(out)
