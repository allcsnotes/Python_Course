from openpyxl import load_workbook

ws = load_workbook("report.xlsx").active
print(ws["C6"].value, ws["C6"].number_format)
print(ws.freeze_panes, ws.column_dimensions["A"].width)
ws = load_workbook("report.xlsx", data_only=True).active
print(ws["C6"].value)
