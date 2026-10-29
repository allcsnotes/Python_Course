from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="1A7CFF")
RUPEES = '"Rs "#,##0'


def style_header(row):
    for cell in row:
        cell.font, cell.fill = HEAD, FILL


def summary_sheet(book, summary, label):
    sheet = book.active
    sheet.title = "Summary"
    sheet["A1"] = f"allcsvideos weekly report {label}"
    sheet["A1"].font = Font(bold=True, size=16)
    sheet.append(["New enrollments", summary.enrollments])
    sheet.append(["Last week", summary.before])
    sheet.append(["Refunds paid", summary.refunds])
    sheet.append(["Refunded (Rs)", summary.refunded])
    sheet.append([])
    sheet.append(["Course", "Enrollments", "Revenue"])
    top = sheet.max_row
    style_header(sheet[top])
    for name, r in summary.per_course.iterrows():
        sheet.append([name, int(r.enrollments), int(r.revenue)])
    last = sheet.max_row
    sheet.append(
        [
            "Total",
            f"=SUM(B{top + 1}:B{last})",
            f"=SUM(C{top + 1}:C{last})",
        ]
    )
    for row in sheet.iter_rows(
        min_row=top + 1, min_col=3, max_col=3
    ):
        row[0].number_format = RUPEES
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["C"].width = 14


def data_sheet(book, summary):
    sheet = book.create_sheet("Enrollments")
    sheet.append(
        ["Student", "Course", "Amount", "Enrolled (India)"]
    )
    style_header(sheet[1])
    for r in summary.rows.itertuples():
        when = r.enrolled_at.replace(tzinfo=None)
        sheet.append([r.student_id, r.course, r.amount, when])
    sheet.freeze_panes = "A2"
    sheet.column_dimensions["D"].width = 22


def build_workbook(summary, label, path):
    book = Workbook()
    summary_sheet(book, summary, label)
    data_sheet(book, summary)
    book.save(path)
