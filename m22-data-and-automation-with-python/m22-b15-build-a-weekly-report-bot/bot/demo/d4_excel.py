from datetime import date, timedelta

import httpx
from openpyxl import load_workbook

from weekly_report.analyse import summarise
from weekly_report.api import fetch_report
from weekly_report.excel import build_workbook
from weekly_report.settings import Settings
from weekly_report.weeks import week_bounds

settings = Settings()
tz = settings.timezone
start, end = week_bounds(date(2026, 10, 19), tz)
with httpx.Client(base_url=settings.api_url) as client:
    data = fetch_report(
        client, settings, start - timedelta(days=7), end
    )
s = summarise(data, start, end, tz)
settings.out_dir.mkdir(exist_ok=True)
path = settings.out_dir / "2026-W43.xlsx"
build_workbook(s, "2026-W43", path)
book = load_workbook(path)
print(book.sheetnames)
sheet = book["Summary"]
for row in sheet.iter_rows(min_row=7, values_only=True):
    print(row)
print(book["Enrollments"].freeze_panes)
